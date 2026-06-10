"""Accounting reports: GL, Trial Balance, P&L, Balance Sheet, VAT, Aging, Excel/PDF export."""
import datetime
from io import BytesIO

from flask import render_template, request, flash, redirect, url_for, send_file
from flask_login import login_required
from sqlalchemy import func

from . import accounting_bp
from ...extensions import db
from ...models.accounting import Account, AccountType, JournalEntry, JournalEntryLine
from ...models.sales import SalesInvoice, Customer, InvoiceType, InvoiceStatus


def _parse_date_range():
    """Parse from_date / to_date from query string; default to current year."""
    today = datetime.date.today()
    default_from = datetime.date(today.year, 1, 1)
    default_to = datetime.date(today.year, 12, 31)
    try:
        from_date = datetime.date.fromisoformat(request.args.get('from_date', ''))
    except (ValueError, TypeError):
        from_date = default_from
    try:
        to_date = datetime.date.fromisoformat(request.args.get('to_date', ''))
    except (ValueError, TypeError):
        to_date = default_to
    return from_date, to_date


def _parse_as_of_date():
    """Parse as_of_date from query string; default to today."""
    today = datetime.date.today()
    try:
        return datetime.date.fromisoformat(request.args.get('as_of_date', ''))
    except (ValueError, TypeError):
        return today


# ---------------------------------------------------------------------------
# 1. General Ledger
# ---------------------------------------------------------------------------

@accounting_bp.route('/reports/ledger')
@login_required
def ledger():
    from_date, to_date = _parse_date_range()

    accounts = Account.query.filter_by(is_active=True).order_by(Account.code).all()
    account_id = request.args.get('account_id', type=int)
    selected_account = None
    lines = []
    opening_balance = 0.0

    if account_id:
        selected_account = db.session.get(Account, account_id)
        if selected_account:
            # Fetch all lines for the account within date range
            lines = (
                db.session.query(JournalEntryLine)
                .join(JournalEntry, JournalEntryLine.entry_id == JournalEntry.id)
                .filter(
                    JournalEntryLine.account_id == account_id,
                    JournalEntry.date >= from_date,
                    JournalEntry.date <= to_date,
                )
                .order_by(JournalEntry.date, JournalEntry.id, JournalEntryLine.id)
                .all()
            )

            # Compute running balance
            balance = opening_balance
            for line in lines:
                balance += float(line.debit or 0) - float(line.credit or 0)
                line._running_balance = balance

    return render_template(
        'accounting/reports/ledger.html',
        accounts=accounts,
        selected_account=selected_account,
        lines=lines,
        from_date=from_date,
        to_date=to_date,
        opening_balance=opening_balance,
    )


# ---------------------------------------------------------------------------
# 2. Trial Balance
# ---------------------------------------------------------------------------

def _get_trial_balance_rows(from_date=None, to_date=None):
    """Return trial balance rows as list of (Account, total_debit, total_credit)."""
    q = (
        db.session.query(
            Account,
            func.coalesce(func.sum(JournalEntryLine.debit), 0).label('total_debit'),
            func.coalesce(func.sum(JournalEntryLine.credit), 0).label('total_credit'),
        )
        .outerjoin(JournalEntryLine, JournalEntryLine.account_id == Account.id)
        .filter(Account.is_active == True)
    )

    if from_date or to_date:
        # Only join entries within the date range when filters are applied
        q = (
            db.session.query(
                Account,
                func.coalesce(func.sum(JournalEntryLine.debit), 0).label('total_debit'),
                func.coalesce(func.sum(JournalEntryLine.credit), 0).label('total_credit'),
            )
            .outerjoin(JournalEntryLine, JournalEntryLine.account_id == Account.id)
            .outerjoin(JournalEntry, JournalEntry.id == JournalEntryLine.entry_id)
            .filter(Account.is_active == True)
        )
        if from_date:
            q = q.filter(
                db.or_(JournalEntry.date == None, JournalEntry.date >= from_date)
            )
        if to_date:
            q = q.filter(
                db.or_(JournalEntry.date == None, JournalEntry.date <= to_date)
            )

    rows = (
        q.group_by(Account.id)
        .having(
            func.coalesce(func.sum(JournalEntryLine.debit), 0) +
            func.coalesce(func.sum(JournalEntryLine.credit), 0) > 0
        )
        .order_by(Account.code)
        .all()
    )
    return rows


@accounting_bp.route('/reports/trial-balance')
@login_required
def trial_balance():
    from_date, to_date = _parse_date_range()
    rows = _get_trial_balance_rows(from_date, to_date)

    # Group by account type
    grouped = {}
    type_order = [AccountType.ASSET, AccountType.LIABILITY, AccountType.EQUITY,
                  AccountType.REVENUE, AccountType.EXPENSE]
    for t in type_order:
        grouped[t] = []

    total_debit = 0.0
    total_credit = 0.0
    for account, td, tc in rows:
        grouped.setdefault(account.type, []).append({
            'account': account,
            'total_debit': float(td),
            'total_credit': float(tc),
            'net': float(td) - float(tc),
        })
        total_debit += float(td)
        total_credit += float(tc)

    return render_template(
        'accounting/reports/trial_balance.html',
        grouped=grouped,
        type_order=type_order,
        AccountType=AccountType,
        total_debit=total_debit,
        total_credit=total_credit,
        from_date=from_date,
        to_date=to_date,
    )


# ---------------------------------------------------------------------------
# 3. Profit & Loss
# ---------------------------------------------------------------------------

@accounting_bp.route('/reports/pl')
@login_required
def profit_loss():
    from_date, to_date = _parse_date_range()

    def _sum_for_type(account_type, from_date, to_date):
        return (
            db.session.query(
                Account,
                func.coalesce(func.sum(JournalEntryLine.debit), 0).label('total_debit'),
                func.coalesce(func.sum(JournalEntryLine.credit), 0).label('total_credit'),
            )
            .outerjoin(JournalEntryLine, JournalEntryLine.account_id == Account.id)
            .outerjoin(JournalEntry, JournalEntry.id == JournalEntryLine.entry_id)
            .filter(
                Account.is_active == True,
                Account.type == account_type,
                db.or_(JournalEntry.date == None, db.and_(
                    JournalEntry.date >= from_date,
                    JournalEntry.date <= to_date,
                )),
            )
            .group_by(Account.id)
            .order_by(Account.code)
            .all()
        )

    revenue_rows = _sum_for_type(AccountType.REVENUE, from_date, to_date)
    expense_rows = _sum_for_type(AccountType.EXPENSE, from_date, to_date)

    revenue_items = []
    total_revenue = 0.0
    for account, td, tc in revenue_rows:
        amount = float(tc) - float(td)  # Revenue: credits - debits
        revenue_items.append({'account': account, 'amount': amount})
        total_revenue += amount

    expense_items = []
    total_expenses = 0.0
    for account, td, tc in expense_rows:
        amount = float(td) - float(tc)  # Expenses: debits - credits
        expense_items.append({'account': account, 'amount': amount})
        total_expenses += amount

    net_profit = total_revenue - total_expenses

    return render_template(
        'accounting/reports/pl.html',
        revenue_items=revenue_items,
        expense_items=expense_items,
        total_revenue=total_revenue,
        total_expenses=total_expenses,
        net_profit=net_profit,
        from_date=from_date,
        to_date=to_date,
    )


# ---------------------------------------------------------------------------
# 4. Balance Sheet
# ---------------------------------------------------------------------------

@accounting_bp.route('/reports/balance-sheet')
@login_required
def balance_sheet():
    as_of_date = _parse_as_of_date()

    def _balance_sheet_rows(account_type):
        return (
            db.session.query(
                Account,
                func.coalesce(func.sum(JournalEntryLine.debit), 0).label('total_debit'),
                func.coalesce(func.sum(JournalEntryLine.credit), 0).label('total_credit'),
            )
            .outerjoin(JournalEntryLine, JournalEntryLine.account_id == Account.id)
            .outerjoin(JournalEntry, JournalEntry.id == JournalEntryLine.entry_id)
            .filter(
                Account.is_active == True,
                Account.type == account_type,
                db.or_(JournalEntry.date == None, JournalEntry.date <= as_of_date),
            )
            .group_by(Account.id)
            .order_by(Account.code)
            .all()
        )

    asset_rows = _balance_sheet_rows(AccountType.ASSET)
    liability_rows = _balance_sheet_rows(AccountType.LIABILITY)
    equity_rows = _balance_sheet_rows(AccountType.EQUITY)

    def _make_items(rows, calc_fn):
        items = []
        total = 0.0
        for account, td, tc in rows:
            amount = calc_fn(float(td), float(tc))
            items.append({'account': account, 'amount': amount})
            total += amount
        return items, total

    asset_items, total_assets = _make_items(asset_rows, lambda d, c: d - c)
    liability_items, total_liabilities = _make_items(liability_rows, lambda d, c: c - d)
    equity_items, total_equity = _make_items(equity_rows, lambda d, c: c - d)

    return render_template(
        'accounting/reports/balance_sheet.html',
        asset_items=asset_items,
        liability_items=liability_items,
        equity_items=equity_items,
        total_assets=total_assets,
        total_liabilities=total_liabilities,
        total_equity=total_equity,
        total_liab_equity=total_liabilities + total_equity,
        as_of_date=as_of_date,
    )


# ---------------------------------------------------------------------------
# 5. VAT Report
# ---------------------------------------------------------------------------

@accounting_bp.route('/reports/vat')
@login_required
def vat_report():
    from_date, to_date = _parse_date_range()

    def _account_sum(code, debit_or_credit, from_date, to_date):
        acc = Account.query.filter_by(code=code).first()
        if not acc:
            return 0.0, None
        col = JournalEntryLine.credit if debit_or_credit == 'credit' else JournalEntryLine.debit
        result = (
            db.session.query(func.coalesce(func.sum(col), 0))
            .join(JournalEntry, JournalEntry.id == JournalEntryLine.entry_id)
            .filter(
                JournalEntryLine.account_id == acc.id,
                JournalEntry.date >= from_date,
                JournalEntry.date <= to_date,
            )
            .scalar()
        )
        return float(result or 0), acc

    output_tax, output_acc = _account_sum('2102', 'credit', from_date, to_date)
    input_tax, input_acc = _account_sum('1105', 'debit', from_date, to_date)
    net_vat = output_tax - input_tax

    return render_template(
        'accounting/reports/vat.html',
        output_tax=output_tax,
        input_tax=input_tax,
        net_vat=net_vat,
        output_acc=output_acc,
        input_acc=input_acc,
        from_date=from_date,
        to_date=to_date,
    )


# ---------------------------------------------------------------------------
# 6. Customer Aging Report
# ---------------------------------------------------------------------------

@accounting_bp.route('/reports/aging')
@login_required
def aging_report():
    today = datetime.date.today()

    # Get all confirmed credit invoices that are outstanding
    invoices = (
        db.session.query(SalesInvoice, Customer)
        .join(Customer, Customer.id == SalesInvoice.customer_id)
        .filter(
            SalesInvoice.type == InvoiceType.CREDIT,
            SalesInvoice.status.in_([InvoiceStatus.CONFIRMED, InvoiceStatus.PARTIAL]),
        )
        .order_by(SalesInvoice.date)
        .all()
    )

    buckets = {
        '0_30': [],
        '31_60': [],
        '61_90': [],
        '90_plus': [],
    }
    totals = {'0_30': 0.0, '31_60': 0.0, '61_90': 0.0, '90_plus': 0.0}

    for inv, customer in invoices:
        outstanding = float(inv.grand_total or 0) - float(inv.amount_paid or 0)
        if outstanding <= 0:
            continue
        inv_date = inv.date if isinstance(inv.date, datetime.date) else inv.date
        days_old = (today - inv_date).days
        row = {
            'invoice': inv,
            'customer': customer,
            'outstanding': outstanding,
            'days_old': days_old,
        }
        if days_old <= 30:
            buckets['0_30'].append(row)
            totals['0_30'] += outstanding
        elif days_old <= 60:
            buckets['31_60'].append(row)
            totals['31_60'] += outstanding
        elif days_old <= 90:
            buckets['61_90'].append(row)
            totals['61_90'] += outstanding
        else:
            buckets['90_plus'].append(row)
            totals['90_plus'] += outstanding

    grand_total = sum(totals.values())

    return render_template(
        'accounting/reports/aging.html',
        buckets=buckets,
        totals=totals,
        grand_total=grand_total,
        today=today,
    )


# ---------------------------------------------------------------------------
# 7. Excel Export — Trial Balance
# ---------------------------------------------------------------------------

@accounting_bp.route('/reports/trial-balance/export/excel')
@login_required
def trial_balance_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        flash('مكتبة openpyxl غير مثبتة', 'danger')
        return redirect(url_for('accounting.trial_balance'))

    from_date, to_date = _parse_date_range()
    rows = _get_trial_balance_rows(from_date, to_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'ميزان المراجعة'

    # Header row
    headers = ['كود الحساب', 'اسم الحساب', 'نوع الحساب', 'إجمالي المدين', 'إجمالي الدائن', 'الرصيد']
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1F4E79')
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    # Data rows
    total_debit = 0.0
    total_credit = 0.0
    for row_idx, (account, td, tc) in enumerate(rows, start=2):
        td_f = float(td)
        tc_f = float(tc)
        net = td_f - tc_f
        total_debit += td_f
        total_credit += tc_f
        ws.cell(row=row_idx, column=1, value=account.code)
        ws.cell(row=row_idx, column=2, value=account.name_ar)
        ws.cell(row=row_idx, column=3, value=AccountType.LABELS.get(account.type, account.type))
        ws.cell(row=row_idx, column=4, value=td_f)
        ws.cell(row=row_idx, column=5, value=tc_f)
        ws.cell(row=row_idx, column=6, value=net)

    # Totals row
    total_row = len(rows) + 2
    ws.cell(row=total_row, column=1, value='الإجمالي').font = Font(bold=True)
    ws.cell(row=total_row, column=4, value=total_debit).font = Font(bold=True)
    ws.cell(row=total_row, column=5, value=total_credit).font = Font(bold=True)
    ws.cell(row=total_row, column=6, value=total_debit - total_credit).font = Font(bold=True)

    # Column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 18

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f'trial_balance_{from_date}_{to_date}.xlsx'
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


# ---------------------------------------------------------------------------
# 8. PDF Export — Trial Balance
# ---------------------------------------------------------------------------

@accounting_bp.route('/reports/trial-balance/export/pdf')
@login_required
def trial_balance_pdf():
    try:
        from weasyprint import HTML
    except ImportError:
        flash('مكتبة WeasyPrint غير مثبتة — لا يمكن تصدير PDF', 'warning')
        from flask import make_response
        resp = make_response(redirect(url_for('accounting.trial_balance')))
        resp.status_code = 501
        return resp

    from_date, to_date = _parse_date_range()
    rows = _get_trial_balance_rows(from_date, to_date)

    grouped = {}
    type_order = [AccountType.ASSET, AccountType.LIABILITY, AccountType.EQUITY,
                  AccountType.REVENUE, AccountType.EXPENSE]
    for t in type_order:
        grouped[t] = []

    total_debit = 0.0
    total_credit = 0.0
    for account, td, tc in rows:
        grouped.setdefault(account.type, []).append({
            'account': account,
            'total_debit': float(td),
            'total_credit': float(tc),
            'net': float(td) - float(tc),
        })
        total_debit += float(td)
        total_credit += float(tc)

    html_content = render_template(
        'accounting/reports/trial_balance_pdf.html',
        grouped=grouped,
        type_order=type_order,
        AccountType=AccountType,
        total_debit=total_debit,
        total_credit=total_credit,
        from_date=from_date,
        to_date=to_date,
    )

    pdf_bytes = HTML(string=html_content).write_pdf()
    output = BytesIO(pdf_bytes)
    output.seek(0)

    filename = f'trial_balance_{from_date}_{to_date}.pdf'
    return send_file(
        output,
        mimetype='application/pdf',
        as_attachment=True,
        download_name=filename,
    )
